"""
Lightweight page count estimation from file bytes before upload.
Used for upload limit enforcement (BILLING_SPEC §5.3).
On any error, returns 1 (conservative — allow upload but minimal estimate).
"""

import logging
from typing import Optional

logger = logging.getLogger(__name__)


def estimate_page_count_from_file(file_content: bytes, filename: str) -> int:
    """
    Estimate page count from file bytes. Used before upload to enforce usage limits.
    Returns 1 on error or for unknown formats (conservative).
    """
    if not file_content or len(file_content) < 10:
        return 1

    name_lower = (filename or "").lower()

    try:
        if name_lower.endswith(".pdf"):
            return _pdf_page_count(file_content)
        if name_lower.endswith(".docx") or name_lower.endswith(".doc"):
            return _docx_page_count(file_content)
        if name_lower.endswith(".xlsx") or name_lower.endswith(".xls"):
            return _excel_page_count(file_content)
        if name_lower.endswith(".csv"):
            return _csv_page_count(file_content)
    except Exception as e:
        logger.warning("Page count estimation failed for %s: %s", filename[:50], e)

    return 1


def _pdf_page_count(file_content: bytes) -> int:
    """PDF: use PyMuPDF for accurate count."""
    try:
        from backend.tasks import get_pdf_page_count_from_bytes

        count = get_pdf_page_count_from_bytes(file_content)
        return max(1, count) if count is not None else 1
    except Exception:
        return 1


def _docx_page_count(file_content: bytes) -> int:
    """DOCX: estimate from text length (~3000 chars per page)."""
    try:
        from docx import Document
        from io import BytesIO

        doc = Document(BytesIO(file_content))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text
        return max(1, len(full_text) // 3000)
    except Exception:
        return 1


def _excel_page_count(file_content: bytes) -> int:
    """Excel: estimate from sheet count (2 pages per sheet). Returns 1 if openpyxl not installed."""
    try:
        import openpyxl  # type: ignore  # optional dependency
        from io import BytesIO

        wb = openpyxl.load_workbook(BytesIO(file_content), read_only=True)
        try:
            sheet_count = len(wb.sheetnames)
            return max(1, sheet_count * 2)
        finally:
            wb.close()
    except (ImportError, Exception):
        return 1


def _csv_page_count(file_content: bytes) -> int:
    """CSV: estimate from row count (~50 rows per page)."""
    try:
        text = file_content.decode("utf-8", errors="ignore")
        lines = [l for l in text.splitlines() if l.strip()]
        return max(1, len(lines) // 50)
    except Exception:
        return 1
